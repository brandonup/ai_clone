"""
Custom RAG implementation with Small-to-Big chunking and sliding window
using Qdrant for vector storage and retrieval
"""
import os
import logging
import json
import re
import uuid
import io # Added for handling bytes with python-docx
from typing import List, Dict, Tuple, Any, Optional, Union

# Set up logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Added imports for text extraction
try:
    from pypdf import PdfReader
except ImportError:
    logger.warning("pypdf not installed. PDF processing will not be available.")
    PdfReader = None
try:
    import docx
except ImportError:
    logger.warning("python-docx not installed. DOCX processing will not be available.")
    docx = None

import tiktoken
from qdrant_client import QdrantClient
from qdrant_client.http import models
from qdrant_client.http.models import Distance, VectorParams, PointStruct
from qdrant_client.models import Filter, FieldCondition, MatchValue
import openai

# Constants
LARGE_CHUNK_SIZE = 512  # tokens
SMALL_CHUNK_SIZE = 175  # tokens
LARGE_CHUNK_OVERLAP = 100  # tokens
SMALL_CHUNK_OVERLAP = 50  # tokens
COLLECTION_NAME = "ai_coach_docs"
EMBEDDING_DIMENSION = 1536  # OpenAI embedding dimension

class DocumentProcessor:
    """
    Process documents with Small-to-Big chunking and sliding window
    """
    
    def __init__(self, 
                 large_chunk_size: int = LARGE_CHUNK_SIZE,
                 small_chunk_size: int = SMALL_CHUNK_SIZE,
                 large_chunk_overlap: int = LARGE_CHUNK_OVERLAP,
                 small_chunk_overlap: int = SMALL_CHUNK_OVERLAP):
        """
        Initialize the document processor
        
        Args:
            large_chunk_size: Size of large chunks in tokens
            small_chunk_size: Size of small chunks in tokens
            large_chunk_overlap: Overlap between large chunks in tokens
            small_chunk_overlap: Overlap between small chunks in tokens
        """
        self.large_chunk_size = large_chunk_size
        self.small_chunk_size = small_chunk_size
        self.large_chunk_overlap = large_chunk_overlap
        self.small_chunk_overlap = small_chunk_overlap
        
        # Initialize tokenizer
        try:
            self.tokenizer = tiktoken.encoding_for_model("gpt-3.5-turbo")
            logger.info("Using tiktoken for token counting")
        except:
            logger.warning("Tiktoken not available, using character-based token estimation")
            self.tokenizer = None
    
    def count_tokens(self, text: str) -> int:
        """
        Count the number of tokens in a text
        
        Args:
            text: Text to count tokens for
            
        Returns:
            int: Number of tokens
        """
        if self.tokenizer:
            return len(self.tokenizer.encode(text))
        else:
            # Fallback to character-based estimation (1 token ≈ 4 chars)
            return len(text) // 4
    
    def preprocess_document_content(self, content: Union[str, bytes]) -> str:
        """
        Preprocess document content to improve chunking and indexing
        
        Args:
            content: Document content as bytes or string
            
        Returns:
            str: Processed content
        """
        # DIAGNOSTIC LOG: Log initial content type and snippet
        logger.info(f"Preprocessing content of type: {type(content)}")
        content_snippet = content[:200] if isinstance(content, str) else content[:200].decode('utf-8', errors='replace')
        logger.info(f"Preprocessing content snippet (first 200 chars/bytes): {content_snippet}")

        # Convert to string if bytes
        if isinstance(content, bytes):
            try:
                text = content.decode('utf-8', errors='strict') # Try strict decoding first
                logger.info("Successfully decoded content as UTF-8 (strict)")
            except UnicodeDecodeError:
                logger.warning("UTF-8 strict decoding failed, falling back to 'replace' errors")
                text = content.decode('utf-8', errors='replace')
        else:
            text = content
        
        # Fix common OCR/PDF extraction issues
        text = text.replace("ﬂ", "fl")  # Replace ligatures
        text = text.replace("ﬁ", "fi")
        
        # Normalize line breaks (replace multiple line breaks with single line break)
        text = re.sub(r'\n{2,}', '\n', text)
        
        # Remove excessive whitespace
        text = re.sub(r'\s+', ' ', text)
        
        # Ensure paragraphs are properly separated
        text = re.sub(r'(\.\s+)([A-Z])', r'\1\n\2', text)
        # DIAGNOSTIC LOG: Log snippet after preprocessing
        logger.info(f"Post-preprocessing snippet (first 200 chars): {text[:200]}")
        return text

    def create_chunks_with_sliding_window(self,
                                          text: str, 
                                          chunk_size: int, 
                                          chunk_overlap: int) -> List[str]:
        """
        Create chunks with sliding window
        
        Args:
            text: Text to chunk
            chunk_size: Size of each chunk in tokens
            chunk_overlap: Overlap between chunks in tokens
            
        Returns:
            List[str]: List of chunks
        """
        if not text:
            return []
        
        # Split text into tokens
        if self.tokenizer:
            tokens = self.tokenizer.encode(text)
            
            # Create chunks with sliding window
            chunks = []
            i = 0
            while i < len(tokens):
                # Get chunk tokens
                chunk_end = min(i + chunk_size, len(tokens))
                chunk_tokens = tokens[i:chunk_end]
                
                # Convert back to text
                chunk_text = self.tokenizer.decode(chunk_tokens)
                chunks.append(chunk_text)
                
                # Slide window
                i += chunk_size - chunk_overlap
                
                # Break if we've reached the end
                if i >= len(tokens):
                    break
            
            return chunks
        else:
            # Fallback to character-based chunking
            # Estimate characters per token (4 chars ≈ 1 token)
            char_chunk_size = chunk_size * 4
            char_chunk_overlap = chunk_overlap * 4
            
            # Create chunks with sliding window
            chunks = []
            i = 0
            while i < len(text):
                # Get chunk
                chunk_end = min(i + char_chunk_size, len(text))
                chunk = text[i:chunk_end]
                chunks.append(chunk)
                
                # Slide window
                i += char_chunk_size - char_chunk_overlap
                
                # Break if we've reached the end
                if i >= len(text):
                    break
            
            return chunks
    
    def process_document(self, 
                         content: Union[str, bytes], 
                         metadata: Dict[str, Any] = None) -> Dict[str, Any]:
        """
        Process a document with Small-to-Big chunking and sliding window
        
        Args:
            content: Document content as bytes or string
            metadata: Document metadata
            
        Returns:
            Dict: Processed document with chunks
        """
        if metadata is None:
            metadata = {}
        
        # Preprocess document content
        text = self.preprocess_document_content(content)
        
        # Create large chunks with sliding window
        large_chunks = self.create_chunks_with_sliding_window(
            text, 
            self.large_chunk_size, 
            self.large_chunk_overlap
        )
        # DIAGNOSTIC LOG: Log first large chunk if available
        if large_chunks:
            logger.info(f"First large chunk created (first 200 chars): {large_chunks[0][:200]}")
        else:
            logger.info("No large chunks created.")

        # Process large chunks
        processed_large_chunks = []
        all_small_chunks = []
        
        for i, large_chunk in enumerate(large_chunks):
            # Create large chunk ID
            large_chunk_id = f"{metadata.get('doc_id', 'doc')}_{i}"
            
            # Create large chunk metadata
            large_chunk_metadata = {
                "chunk_id": large_chunk_id,
                "doc_id": metadata.get("doc_id", ""),
                "doc_title": metadata.get("title", ""),
                "chunk_type": "large",
                "chunk_index": i,
            }
            
            # Add large chunk
            processed_large_chunks.append({
                "text": large_chunk,
                "metadata": large_chunk_metadata
            })
            
            # Create small chunks with sliding window
            small_chunks = self.create_chunks_with_sliding_window(
                large_chunk, 
                self.small_chunk_size, 
                self.small_chunk_overlap
            )
            
            # Process small chunks
            for j, small_chunk in enumerate(small_chunks):
                # Create small chunk ID
                small_chunk_id = f"{large_chunk_id}_{j}"
                
                # Create small chunk metadata
                small_chunk_metadata = {
                    "chunk_id": small_chunk_id,
                    "doc_id": metadata.get("doc_id", ""),
                    "doc_title": metadata.get("title", ""),
                    "chunk_type": "small",
                    "chunk_index": j,
                    "parent_id": large_chunk_id,
                    "parent_text": large_chunk, # Storing potentially large text here
                }

                # Add small chunk
                all_small_chunks.append({
                    "text": small_chunk,
                    "metadata": small_chunk_metadata
                })

            # DIAGNOSTIC LOG: Log first small chunk of the first large chunk
            if i == 0 and small_chunks:
                 logger.info(f"First small chunk created (from first large chunk, first 200 chars): {small_chunks[0][:200]}")
            elif i == 0:
                 logger.info("No small chunks created from the first large chunk.")

        # Return processed document
        return {
            "doc_id": metadata.get("doc_id", ""),
            "title": metadata.get("title", ""),
            "large_chunks": processed_large_chunks,
            "small_chunks": all_small_chunks,
        }

class QdrantStore:
    """
    Store and retrieve document chunks using Qdrant
    """
    
    def __init__(self, 
                 url: Optional[str] = None,
                 api_key: Optional[str] = None,
                 collection_name: str = COLLECTION_NAME):
        """
        Initialize the Qdrant store
        
        Args:
            url: Qdrant server URL
            api_key: Qdrant API key
            collection_name: Qdrant collection name
        """
        self.collection_name = collection_name
        
        # Initialize Qdrant client
        if url:
            self.client = QdrantClient(url=url, api_key=api_key)
        else:
            # Use in-memory Qdrant for testing
            self.client = QdrantClient(":memory:")
        
        # Initialize OpenAI client
        openai_api_key = os.getenv("OPENAI_API_KEY")
        if not openai_api_key:
            raise ValueError("OPENAI_API_KEY environment variable not set")
        
        self.openai_client = openai.OpenAI(api_key=openai_api_key)
        
        # Ensure collection exists
        self._ensure_collection_exists()
    
    def _ensure_collection_exists(self):
        """
        Ensure the Qdrant collection exists
        """
        # Check if collection exists
        collections = self.client.get_collections().collections
        collection_names = [collection.name for collection in collections]
        
        if self.collection_name not in collection_names:
            # Create collection
            self.client.create_collection(
                collection_name=self.collection_name,
                vectors_config=VectorParams(
                    size=EMBEDDING_DIMENSION,
                    distance=Distance.COSINE
                )
            )
            logger.info(f"Created Qdrant collection: {self.collection_name}")
    
    def generate_embeddings(self, texts: List[str]) -> List[List[float]]:
        """
        Generate embeddings for texts using OpenAI
        
        Args:
            texts: List of texts to generate embeddings for
            
        Returns:
            List[List[float]]: List of embeddings
        """
        # Handle rate limits by processing in batches
        batch_size = 100  # Process 100 texts at a time to avoid rate limits
        all_embeddings = []
        
        for i in range(0, len(texts), batch_size):
            batch_texts = texts[i:i+batch_size]
            
            try:
                # Generate embeddings for this batch
                logger.info(f"Generating embeddings for batch {i//batch_size + 1}/{(len(texts)-1)//batch_size + 1}")
                response = self.openai_client.embeddings.create(
                    model="text-embedding-ada-002",
                    input=batch_texts
                )
                
                # Extract embeddings
                batch_embeddings = [item.embedding for item in response.data]
                all_embeddings.extend(batch_embeddings)
                
                # Sleep briefly to avoid hitting rate limits
                if i + batch_size < len(texts):
                    import time
                    time.sleep(0.5)  # 500ms pause between batches
                    
            except Exception as e:
                logger.error(f"Error generating embeddings for batch {i//batch_size + 1}: {str(e)}")
                # If we hit a rate limit, wait longer and retry
                if "rate_limit_exceeded" in str(e):
                    logger.info("Rate limit exceeded, waiting 60 seconds before retrying...")
                    import time
                    time.sleep(60)
                    
                    # Retry with a smaller batch
                    smaller_batch_size = batch_size // 2
                    for j in range(i, min(i+batch_size, len(texts)), smaller_batch_size):
                        sub_batch = texts[j:j+smaller_batch_size]
                        try:
                            logger.info(f"Retrying with smaller batch of {len(sub_batch)} texts")
                            response = self.openai_client.embeddings.create(
                                model="text-embedding-ada-002",
                                input=sub_batch
                            )
                            sub_batch_embeddings = [item.embedding for item in response.data]
                            all_embeddings.extend(sub_batch_embeddings)
                            time.sleep(1)  # Longer pause between retries
                        except Exception as sub_e:
                            logger.error(f"Error in retry batch: {str(sub_e)}")
                            # For each text that fails, use a zero vector as placeholder
                            # This allows processing to continue even if some embeddings fail
                            for _ in range(len(sub_batch)):
                                all_embeddings.append([0.0] * EMBEDDING_DIMENSION)
                else:
                    # For non-rate-limit errors, use zero vectors as placeholders
                    for _ in range(len(batch_texts)):
                        all_embeddings.append([0.0] * EMBEDDING_DIMENSION)
        
        return all_embeddings
    
    def store_document(self, processed_doc: Dict[str, Any]) -> None:
        """
        Store a processed document in Qdrant
        
        Args:
            processed_doc: Processed document with chunks
        """
        # Extract small chunks
        small_chunks = processed_doc["small_chunks"]
        
        if not small_chunks:
            logger.warning(f"No small chunks to store for document: {processed_doc['doc_id']}")
            return
        
        # Extract texts and metadata
        texts = [chunk["text"] for chunk in small_chunks]
        metadatas = [chunk["metadata"] for chunk in small_chunks]
        
        # Generate embeddings
        embeddings = self.generate_embeddings(texts)
        
        # Create points
        points = []
        for i, (text, metadata, embedding) in enumerate(zip(texts, metadatas, embeddings)):
            # Create point ID
            point_id = str(uuid.uuid4())
            
            # Create point
            point = PointStruct(
                id=point_id,
                vector=embedding,
                payload={
                    "text": text,
                    "chunk_id": metadata["chunk_id"],
                    "doc_id": metadata["doc_id"],
                    "doc_title": metadata["doc_title"],
                    "chunk_type": metadata["chunk_type"],
                    "chunk_index": metadata["chunk_index"],
                    "parent_id": metadata["parent_id"],
                    "parent_text": metadata["parent_text"],
                }
            )
            
            points.append(point)
        
        # Upsert points in batches to avoid payload size limits
        batch_size = 100  # Process 100 points at a time to avoid payload size limits
        total_points = len(points)
        
        logger.info(f"Upserting {total_points} points in batches of {batch_size}")
        
        for i in range(0, total_points, batch_size):
            batch_end = min(i + batch_size, total_points)
            batch_points = points[i:batch_end]
            
            try:
                logger.info(f"Upserting batch {i//batch_size + 1}/{(total_points-1)//batch_size + 1} ({len(batch_points)} points)")
                self.client.upsert(
                    collection_name=self.collection_name,
                    points=batch_points
                )
                
                # Sleep briefly between batches to avoid rate limits
                if i + batch_size < total_points:
                    import time
                    time.sleep(0.5)  # 500ms pause between batches
                    
            except Exception as e:
                logger.error(f"Error upserting batch {i//batch_size + 1}: {str(e)}")
                # If we hit a rate limit or size limit, try with smaller batches
                if "Payload error" in str(e) or "rate_limit_exceeded" in str(e):
                    logger.info("Payload too large or rate limit exceeded, retrying with smaller batches")
                    
                    # Retry with smaller batches
                    smaller_batch_size = batch_size // 2
                    for j in range(i, batch_end, smaller_batch_size):
                        sub_batch_end = min(j + smaller_batch_size, batch_end)
                        sub_batch_points = points[j:sub_batch_end]
                        
                        try:
                            logger.info(f"Upserting smaller batch of {len(sub_batch_points)} points")
                            self.client.upsert(
                                collection_name=self.collection_name,
                                points=sub_batch_points
                            )
                            import time
                            time.sleep(1)  # Longer pause between retries
                        except Exception as sub_e:
                            logger.error(f"Error in retry batch: {str(sub_e)}")
                            # If even smaller batches fail, try one by one
                            logger.info("Trying to upsert points one by one")
                            for point in sub_batch_points:
                                try:
                                    self.client.upsert(
                                        collection_name=self.collection_name,
                                        points=[point]
                                    )
                                    time.sleep(0.2)  # Small pause between individual points
                                except Exception as point_e:
                                    logger.error(f"Failed to upsert point: {str(point_e)}")
                else:
                    # For other errors, log and continue
                    logger.error(f"Unexpected error during upsert: {str(e)}")
        
        logger.info(f"Stored {total_points} small chunks for document: {processed_doc['doc_id']}")
    
    def search(self, 
               query: str, 
               limit: int = 5, 
               filter: Optional[Filter] = None) -> List[Dict[str, Any]]:
        """
        Search for similar chunks
        
        Args:
            query: Query text
            limit: Maximum number of results
            filter: Qdrant filter
            
        Returns:
            List[Dict]: List of search results
        """
        # Generate query embedding
        query_embedding = self.generate_embeddings([query])[0]
        
        # Search for similar chunks
        search_result = self.client.search(
            collection_name=self.collection_name,
            query_vector=query_embedding,
            limit=limit,
            query_filter=filter
        )
        
        # Extract results
        results = []
        for scored_point in search_result:
            # Extract point
            point_id = scored_point.id
            score = scored_point.score
            payload = scored_point.payload
            
            # Create result
            result = {
                "id": point_id,
                "score": score,
                "text": payload["text"],
                "chunk_id": payload["chunk_id"],
                "doc_id": payload["doc_id"],
                "doc_title": payload["doc_title"],
                "chunk_type": payload["chunk_type"],
                "chunk_index": payload["chunk_index"],
                "parent_id": payload["parent_id"],
                "parent_text": payload["parent_text"],
            }
            
            results.append(result)
        
        return results

def ingest_document(file_content: Union[str, bytes], filename: str) -> Dict[str, Any]:
    """
    Ingest a document into the custom RAG system
    
    Args:
        file_content: File content as bytes or string
        filename: Name of the file
        
    Returns:
        Dict: Response with document ID
    """
    logger.info(f"Starting ingestion for filename: {filename}")
    
    # DIAGNOSTIC LOG: Log initial file_content type and size
    content_size = len(file_content) if isinstance(file_content, str) else len(file_content)
    logger.info(f"ingest_document received content of type: {type(file_content)}, size: {content_size} bytes/chars")
    
    try:
        if isinstance(file_content, str):
            content_snippet_ingest = file_content[:200]
        elif isinstance(file_content, bytes):
            try:
                content_snippet_ingest = file_content[:200].decode('utf-8', errors='replace')
            except Exception as decode_e:
                content_snippet_ingest = f"[Binary data, first 20 bytes: {file_content[:20].hex()}]"
        else:
            content_snippet_ingest = f"[Unknown content type: {type(file_content)}]"
            
        logger.info(f"ingest_document received content snippet (first 200 chars/bytes): {content_snippet_ingest}")
    except Exception as log_e:
        logger.error(f"Error logging initial content snippet in ingest_document: {log_e}")

    # --- Text Extraction Logic ---
    extracted_text = None
    file_lower = filename.lower()

    if file_lower.endswith('.pdf') and PdfReader:
        logger.info(f"Processing PDF file: {filename}")
        try:
            if isinstance(file_content, bytes):
                pdf_file = io.BytesIO(file_content)
                reader = PdfReader(pdf_file)
                extracted_text = ""
                for page in reader.pages:
                    extracted_text += page.extract_text() + "\n"
                logger.info(f"Successfully extracted text from PDF: {filename}")
            else:
                logger.error(f"Expected bytes for PDF processing, got {type(file_content)}")
        except Exception as e:
            logger.error(f"Error extracting text from PDF {filename}: {str(e)}")
            # Optionally raise error or return failure

    elif file_lower.endswith('.docx') and docx:
        logger.info(f"Processing DOCX file: {filename}")
        try:
            if isinstance(file_content, bytes):
                doc_file = io.BytesIO(file_content)
                document = docx.Document(doc_file)
                extracted_text = "\n".join([para.text for para in document.paragraphs])
                logger.info(f"Successfully extracted text from DOCX: {filename}")
            else:
                logger.error(f"Expected bytes for DOCX processing, got {type(file_content)}")
        except Exception as e:
            logger.error(f"Error extracting text from DOCX {filename}: {str(e)}")
            # Optionally raise error or return failure

    elif file_lower.endswith('.txt'):
        logger.info(f"Processing TXT file: {filename}")
        if isinstance(file_content, bytes):
            try:
                # Try decoding with UTF-8 first (most common)
                extracted_text = file_content.decode('utf-8', errors='strict')
                logger.info(f"Successfully decoded TXT as UTF-8: {filename}")
            except UnicodeDecodeError:
                logger.warning(f"UTF-8 decoding failed for {filename}, trying latin-1")
                try:
                    # Fallback to latin-1 if UTF-8 fails
                    extracted_text = file_content.decode('latin-1', errors='replace')
                    logger.info(f"Successfully decoded TXT as latin-1: {filename}")
                except Exception as e:
                     logger.error(f"Error decoding TXT {filename} even with fallback: {str(e)}")
        elif isinstance(file_content, str):
             extracted_text = file_content # Assume already decoded if string
        else:
             logger.error(f"Expected bytes or str for TXT processing, got {type(file_content)}")

    else:
        logger.warning(f"Unsupported file type or missing library for: {filename}. Attempting direct decode.")
        if isinstance(file_content, bytes):
             extracted_text = file_content.decode('utf-8', errors='replace') # Fallback attempt
        elif isinstance(file_content, str):
             extracted_text = file_content
        else:
             logger.error(f"Cannot process content type {type(file_content)} for {filename}")

    if extracted_text is None:
         logger.error(f"Failed to extract text from {filename}. Skipping ingestion.")
         # Return an error or empty success indicator
         return {"status": "error", "message": f"Failed to extract text from {filename}"}
    # --- End Text Extraction Logic ---


    # Create document ID
    doc_id = str(uuid.uuid4())

    # Create document metadata
    metadata = {
        "doc_id": doc_id,
        "title": filename,
    }

    # Use the extracted_text for processing
    try:
        # Check if the extracted text is very large
        if len(extracted_text) > 1000000:  # More than ~1MB of text
            logger.info(f"Large document detected ({len(extracted_text)} chars). Processing in sections.")

            # Break into major sections (roughly 500KB each)
            section_size = 500000  # ~500KB per section
            sections = [extracted_text[i:i+section_size] for i in range(0, len(extracted_text), section_size)]
            
            logger.info(f"Document split into {len(sections)} sections for processing")

            # Process each section separately
            processor = DocumentProcessor()

            # Initialize Qdrant store with URL and API key from environment variables
            qdrant_url = os.getenv("QDRANT_URL")
            qdrant_api_key = os.getenv("QDRANT_API_KEY")
            
            if not qdrant_url or not qdrant_api_key:
                logger.error("QDRANT_URL or QDRANT_API_KEY environment variables not set")
                raise ValueError("QDRANT_URL or QDRANT_API_KEY environment variables not set")
            
            store = QdrantStore(url=qdrant_url, api_key=qdrant_api_key)
            
            for i, section in enumerate(sections):
                section_metadata = metadata.copy()
                section_metadata["section"] = i
                
                logger.info(f"Processing section {i+1}/{len(sections)}")
                processed_section = processor.process_document(section, section_metadata)
                
                logger.info(f"Storing section {i+1}/{len(sections)}")
                store.store_document(processed_section)
                
                # Sleep between sections to avoid rate limits
                if i < len(sections) - 1:
                    import time
                    time.sleep(2)  # 2 second pause between sections
        else:
            # Process document normally
            processor = DocumentProcessor()
            processed_doc = processor.process_document(file_content, metadata)
            
            # Initialize Qdrant store with URL and API key from environment variables
            qdrant_url = os.getenv("QDRANT_URL")
            qdrant_api_key = os.getenv("QDRANT_API_KEY")
            
            if not qdrant_url or not qdrant_api_key:
                logger.error("QDRANT_URL or QDRANT_API_KEY environment variables not set")
                raise ValueError("QDRANT_URL or QDRANT_API_KEY environment variables not set")

            # Store document
            store = QdrantStore(url=qdrant_url, api_key=qdrant_api_key)
            # Pass the extracted_text to the processor
            processed_doc = processor.process_document(extracted_text, metadata)
            store.store_document(processed_doc)
    except Exception as e:
        logger.error(f"Error processing document content for {filename}: {str(e)}")
        # Re-raise with more context
        raise Exception(f"Error processing document content for {filename}: {str(e)}")

    # Store a local copy of the *extracted text* for fallback
    try:
        # Use the extracted_text string directly
        doc_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'documents', filename + ".txt") # Save as .txt
        os.makedirs(os.path.dirname(doc_path), exist_ok=True)

        with open(doc_path, 'w', encoding='utf-8') as f: # Ensure writing as UTF-8
            f.write(extracted_text)

        logger.info(f"Stored local copy of extracted text at {doc_path}")
    except Exception as e:
        logger.error(f"Error storing local copy of extracted text for {filename}: {str(e)}")

    # Return response
    return {
        "status": "success",
        "doc_id": doc_id,
    }

def retrieve_chunks(query: str, top_n: int = 1) -> List[str]:
    """
    Retrieve relevant chunks for a given query
    
    Args:
        query: User question
        top_n: Number of top chunks to return
        
    Returns:
        List[str]: List of text chunks
    """
    logger.info(f"Original query: {query}")
    
    # Try multiple query variations to improve retrieval
    queries = [
        query,  # Original query
        f"Find information about: {query}",  # Directive format
        query.replace("?", "").strip(),  # Remove question mark
        " ".join(query.lower().split()),  # Normalize spacing and case
    ]
    
    # Initialize Qdrant store with URL and API key from environment variables
    qdrant_url = os.getenv("QDRANT_URL")
    qdrant_api_key = os.getenv("QDRANT_API_KEY")
    
    if not qdrant_url or not qdrant_api_key:
        logger.error("QDRANT_URL or QDRANT_API_KEY environment variables not set")
        return []
    
    logger.info(f"Connecting to Qdrant at {qdrant_url}")
    store = QdrantStore(url=qdrant_url, api_key=qdrant_api_key)
    
    # Search for each query variation
    all_results = []
    
    for q in queries:
        if not q.strip():  # Skip empty queries
            continue
            
        logger.info(f"Trying query variation: {q}")
        
        # Search for similar chunks
        results = store.search(q, limit=10)
        
        if results:
            logger.info(f"Query '{q}' returned {len(results)} chunks")
            all_results.extend(results)
        else:
            logger.info(f"Query '{q}' returned no chunks")
    
    # If no results, log a warning and return an empty list
    if not all_results:
        logger.warning("No chunks returned from vector search, returning empty list")
        return []
    
    # Remove duplicates while preserving order
    unique_results = []
    seen_parent_ids = set()
    
    for result in all_results:
        parent_id = result["parent_id"]
        
        if parent_id not in seen_parent_ids:
            seen_parent_ids.add(parent_id)
            unique_results.append(result)
    
    # Sort by score (highest first)
    unique_results.sort(key=lambda x: x["score"], reverse=True)
    
    # Get parent chunks for top results
    parent_chunks = []
    
    for result in unique_results[:top_n]:
        parent_text = result["parent_text"]
        parent_chunks.append(parent_text)
    
    # Log the final selected chunks
    logger.info(f"Selected {len(parent_chunks)} parent chunks")
    for i, text in enumerate(parent_chunks):
        logger.info(f"Parent chunk {i+1}: {text[:100]}...")
    
    # Return the parent chunks directly without using ragie_utils.sanitize_text
    logger.info(f"Returning {len(parent_chunks)} parent chunks")
    for i, text in enumerate(parent_chunks):
        logger.info(f"Parent chunk {i+1}: {text[:100]}...")
    return parent_chunks
    logger.info(f"Returning {len(parent_chunks)} parent chunks")
    for i, text in enumerate(parent_chunks):
        logger.info(f"Parent chunk {i+1}: {text[:100]}...")
    return parent_chunks
